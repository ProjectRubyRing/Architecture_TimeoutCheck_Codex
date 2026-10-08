"""Check artifact contracts, references, initial caches and failure scenarios."""

import json
import re
import sys
import unittest
from pathlib import Path

import openpyxl
from PIL import Image, ImageChops

from guide_budget import CHECKS, PARAMETERS, evaluate, statuses, values
from guide_examples import EXAMPLES, TEST_PLAN
from guide_routes import ROUTES
from guide_settings import SETTINGS
from guide_sources import SOURCES
from guide_topics import CHAPTERS

ROOT = Path(__file__).resolve().parent.parent
STEM = 'AWS_Timeout_Design_Guide'


class ArtifactTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.book = openpyxl.load_workbook(ROOT/f'{STEM}.xlsx', data_only=False)
        cls.cached = openpyxl.load_workbook(ROOT/f'{STEM}.xlsx', data_only=True)
        cls.markdown = (ROOT/f'{STEM}.md').read_text(encoding='utf-8')

    @classmethod
    def tearDownClass(cls):
        cls.book.close()
        cls.cached.close()

    def test_twenty_one_sheets_and_reading_order(self):
        expected = ['00_読み方','01_設定一覧','02_整合性計算']
        expected += [r['sheet'] for r in ROUTES]
        expected += [c['sheet'] for c in CHAPTERS if not c['sheet'].startswith('19_')]
        expected += ['17_設定と実装例','18_検証計画','19_用語集','20_公式出典']
        self.assertEqual(self.book.sheetnames, expected)
        self.assertEqual(len(expected), 21)

    def test_all_populated_cells_use_meiryo_ui_and_monochrome(self):
        for ws in self.book:
            for row in ws:
                for cell in row:
                    if cell.value is None:
                        continue
                    self.assertEqual(cell.font.name, 'Meiryo UI', f'{ws.title}!{cell.coordinate}')
                    for color in [cell.font.color, cell.fill.fgColor]:
                        if color is not None and color.type == 'rgb':
                            rgb = color.rgb[-6:]
                            self.assertEqual(rgb[0:2], rgb[2:4])
                            self.assertEqual(rgb[2:4], rgb[4:6])
            self.assertIsNotNone(ws.print_area)
            self.assertIsNotNone(ws.freeze_panes)

    def test_complete_initial_formula_caches(self):
        formula_count = 0
        for ws in self.book:
            for row in ws:
                for cell in row:
                    if cell.data_type != 'f':
                        continue
                    formula_count += 1
                    cached = self.cached[ws.title][cell.coordinate]
                    self.assertNotEqual(cached.data_type, 'e', f'{ws.title}!{cell.coordinate}')
                    self.assertIsNotNone(cached.value)
                    if cell.column == 6:
                        self.assertEqual(cached.value, 'OK', cell.coordinate)
        self.assertEqual(formula_count, len(PARAMETERS)+2*len(CHECKS)+1)

    def test_named_input_cells_and_validation(self):
        for p in PARAMETERS:
            self.assertIn(p['name'], self.book.defined_names)
            destinations = list(self.book.defined_names[p['name']].destinations)
            self.assertEqual(len(destinations), 1)
            sheet, coordinate = destinations[0]
            self.assertEqual(sheet, '02_整合性計算')
            cell = self.book[sheet][coordinate]
            self.assertEqual(cell.value, p['value'])
            self.assertTrue(any(cell.coordinate in dv for dv in self.book[sheet].data_validations.dataValidation))

    def test_initial_cached_margins(self):
        calc = self.cached['02_整合性計算']
        found = {}
        for row in calc:
            if row[0].value and isinstance(row[0].value, str):
                cid = row[0].value.split('\n')[-1]
                if re.fullmatch(r'B\d{2}', cid):
                    found[cid] = row[2].value
        self.assertEqual(set(found), {c['id'] for c in CHECKS})
        for c in CHECKS:
            self.assertAlmostEqual(found[c['id']], evaluate(c['expression'], values()), places=6)
        self.assertEqual(found['B02'], 2.5)
        self.assertEqual(found['B15'], 5)
        self.assertEqual(found['B27'], 0)
        self.assertEqual(found['B38'], 9)

    def test_internal_links_have_real_destinations(self):
        internal_count = 0
        for ws in self.book:
            for row in ws:
                for cell in row:
                    h = cell.hyperlink
                    if h is None or not h.location:
                        continue
                    internal_count += 1
                    m = re.fullmatch(r"(?:'([^']+)'|([^!]+))!(\$?[A-Z]+\$?\d+)", h.location)
                    self.assertIsNotNone(m, h.location)
                    target, coordinate = m.group(1) or m.group(2), m.group(3)
                    self.assertIn(target, self.book.sheetnames)
                    self.assertIsNotNone(self.book[target][coordinate].value, h.location)
        self.assertGreater(internal_count, 250)

    def test_markdown_and_sources_are_complete(self):
        self.assertNotIn('\ufffd', self.markdown)
        self.assertNotIn('\x00', self.markdown)
        for i in range(21):
            self.assertEqual(self.markdown.count(f'<a id="s{i:02d}"></a>'), 1)
        for s in SETTINGS:
            self.assertEqual(self.markdown.count(f'<a id="setting-{s["id"]}"></a>'), 1)
        links = {cell.hyperlink.target for row in self.book['20_公式出典'] for cell in row if cell.hyperlink}
        for sid, title, publisher, url, note in SOURCES:
            self.assertIn(url, links)
            self.assertIn(url, self.markdown)
        self.assertEqual(len([url for url in links if url and url.startswith('https://')]), len(SOURCES))
        for test in TEST_PLAN:
            self.assertIn(test[0], self.markdown)

    def test_six_embedded_monochrome_diagrams(self):
        for route in ROUTES:
            path = ROOT/'diagrams'/f'{route["id"]}_route.png'
            self.assertIn(f'(diagrams/{path.name})', self.markdown)
            self.assertEqual(len(self.book[route['sheet']]._images), 1)
            with Image.open(path) as img:
                self.assertEqual(img.size, (1800,1160))
                red, green, blue = img.convert('RGB').split()
                self.assertIsNone(ImageChops.difference(red, green).getbbox())
                self.assertIsNone(ImageChops.difference(red, blue).getbbox())

    def test_settings_examples_parse_and_tests_are_not_claimed_run(self):
        for e in EXAMPLES:
            if e['language'] == 'json':
                json.loads(e['code'])
        self.assertEqual(len(TEST_PLAN), 32)
        sheet = self.book['18_検証計画']
        for row in sheet.iter_rows(min_row=6, max_row=5+len(TEST_PLAN)):
            self.assertEqual(row[6].value, '未実施')

    def test_initial_budgets_and_independent_failure_scenarios(self):
        self.assertTrue(all(statuses().values()))
        # These scenarios assert the design contract, independently of Excel rows.
        for overrides, expected in [
            ({'p_sql':10}, {'B06','B07','B08'}),
            ({'p_report_idle':40}, {'B17'}),
            ({'p_visibility':60}, {'B27'}),
            ({'p_lambda_hard':80}, {'B27'}),
            ({'p_window':5}, {'B27'}),
            ({'p_batch':10}, {'B25'}),
            ({'p_gwlb':120,'p_nlb':120}, {'B36'}),
            ({'p_token_refresh':3600}, {'B32'}),
            ({'p_gateway_elapsed':25}, {'B38','B39'}),
            ({'p_stop':30}, {'B43'}),
        ]:
            failed = {cid for cid, ok in statuses(overrides).items() if not ok}
            self.assertTrue(expected <= failed, (overrides, failed))

    def test_tighter_deadline_never_expands_work(self):
        check = next(c for c in CHECKS if c['id']=='B38')
        previous = 20
        for elapsed in range(31):
            result = evaluate(check['expression'], values({'p_gateway_elapsed':elapsed}))
            self.assertLessEqual(result, previous)
            self.assertGreaterEqual(result, 0)
            previous = result

    def test_source_text_has_no_invalid_whitespace_or_encoding(self):
        paths = list((ROOT/'tools').glob('*.py')) + [ROOT/'README.md',ROOT/f'{STEM}.md']
        for path in paths:
            text = path.read_text(encoding='utf-8')
            self.assertNotIn('\ufffd', text, path.name)
            self.assertNotIn('\x00', text, path.name)
            for line, content in enumerate(text.splitlines(), start=1):
                self.assertEqual(content.rstrip(), content, f'{path.name}:{line} trailing whitespace')


if __name__ == '__main__':
    suite = unittest.defaultTestLoader.loadTestsFromTestCase(ArtifactTests)
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    report = dict(artifact=f'{STEM}.xlsx', validation_date='2026-10-08',
                  tests_run=result.testsRun, failures=len(result.failures), errors=len(result.errors),
                  result='passed' if result.wasSuccessful() else 'failed',
                  sheets=21, settings=len(SETTINGS), editable_inputs=len(PARAMETERS),
                  budget_conditions=len(CHECKS), sources=len(SOURCES),
                  environment_tests_executed=False,
                  tests=[name for name in dir(ArtifactTests) if name.startswith('test_')])
    directory = ROOT/'validation'
    directory.mkdir(exist_ok=True)
    (directory/'guide_checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    sys.exit(0 if result.wasSuccessful() else 1)
